from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def generate_excel(columns, rows, payroll, service=None):
    """
    Generate a payroll contribution Excel file from query results.
    """

    output_folder = Path("output")
    output_folder.mkdir(exist_ok=True)

    safe_payroll = payroll.replace("/", "_")

    if service:
        filename = f"DPF_{safe_payroll}_{service}.xlsx"
    else:
        filename = f"DPF_{safe_payroll}.xlsx"

    output_path = output_folder / filename

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Payroll Contributions"

    # Write column headings
    for column_index, column_name in enumerate(columns, start=1):
        worksheet.cell(
            row=1,
            column=column_index,
            value=column_name
        )

    # Write returned payroll data
    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )

    # Basic automatic column sizing
    for column_index, column_name in enumerate(columns, start=1):
        column_letter = get_column_letter(column_index)

        max_length = len(str(column_name))

        for cell in worksheet[column_letter]:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[column_letter].width = min(
            max_length + 2,
            40
        )

    workbook.save(output_path)

    return output_path